# -*- coding: utf-8 -*-
"""seeded K=3：种子质心=10钓鱼+10垃圾+10正常，评估+簇构成"""
import json, re
import numpy as np
import openpyxl
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.cluster import KMeans
from sklearn.metrics import adjusted_rand_score, normalized_mutual_info_score, precision_recall_fscore_support
from scipy.sparse import hstack, csr_matrix

XLSX = 'data/法语语料库_最终版_钓鱼邮件319条+垃圾短信1272条.xlsx'
STOP = set("""le la les un une des et ou mais donc or ni car à de du au aux en dans sur sous avec sans par pour que qui quoi dont où ne pas plus moins très tout tous toute toutes ce cet cette ces il elle ils elles on nous vous je tu mon ma mes ton ta tes son sa ses notre votre leur nos vos leur se si y en a est sont était étaient être avoir avait avaient fait faites faire peut peuvent doit doivent comme aussi ainsi alors lorsque quand après avant entre contre chez vers parmi depuis pendant tant selon""".split())
URL_RE = re.compile(r'https?://|www\.|\b[a-z0-9-]+\.(com|net|org|fr|eu|io|ly|xyz|top|club|site|info)\b', re.I)
SHORT_RE = re.compile(r'\b(bit\.ly|t\.co|goo\.gl|tinyurl|shorturl|ow\.ly|is\.gd|buff\.ly|cutt\.ly)\b|https?://[^\s]{1,12}\b', re.I)
MONEY_RE = re.compile(r'[€$£]|\b\d[\d\s.,]{2,}\s?(euros?|€|dollars?)\b', re.I)
BANK = re.compile(r'\b(banque|compte|relevé|identifiant|identif|code|mot de passe|carte|crédit|virement|transaction|sécurit|paypal|amazon|dhl|la poste|imp[ôô]ts|police|gendarmerie|arcep|free|orange|sosh|bouygues)\b', re.I)
THREAT = re.compile(r'\b(p[ée]nalit[ée]|amende|poursuites|pirat[ée]|ransomware|webcam|chantage|dossier|proc[èe]s|prison|virus|infection)\b', re.I)

def clean(t):
    t = re.sub(r'https?://\S+|www\.\S+|\S+@\S+|\b\d[\d\s\-.()/]*\d', ' ', t.lower())
    t = re.sub(r'[^\w\sàâäéèêëîïôöùûüçœ]', ' ', t, flags=re.UNICODE)
    return ' '.join(w for w in t.split() if w not in STOP and len(w) > 2)

def feats(t):
    tl = t.lower()
    return [1 if URL_RE.search(tl) else 0, 1 if SHORT_RE.search(tl) else 0,
            1 if MONEY_RE.search(tl) else 0, 1 if BANK.search(tl) else 0,
            1 if THREAT.search(tl) else 0, len(t)]

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
ph, sm = [], []
ws = wb['钓鱼邮件语料']
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not row[0] or not row[1]: continue
    ph.append(str(row[1]))
ws2 = wb['垃圾短信数据']
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0 or not row[0] or not row[1]: continue
    sm.append({'text': str(row[1]), 'label': int(row[2] if row[2] is not None else 1)})
wb.close()

docs = ph + [s['text'] for s in sm]
true = np.array([1]*len(ph) + [0]*len(sm))
kind = np.array(['钓鱼']*len(ph) + ['垃圾' if s['label']==1 else '正常' for s in sm])
n_ph = len(ph)
spam_idx = [n_ph + i for i, s in enumerate(sm) if s['label'] == 1]
norm_idx = [n_ph + i for i, s in enumerate(sm) if s['label'] == 0]

F = np.array([feats(t) for t in docs], dtype=float); F[:, -1] /= 400.0
vec = TfidfVectorizer(ngram_range=(1,2), max_features=8000, sublinear_tf=True)
T = vec.fit_transform([clean(d) for d in docs])
X = hstack([T, csr_matrix(F)]).tocsr()

def center(rows):
    return np.asarray(X[rows].mean(axis=0)).ravel()

seeds = np.vstack([center(list(range(10))), center(spam_idx[:10]), center(norm_idx[:10])])
km = KMeans(n_clusters=3, init=seeds, n_init=1, random_state=42)
pred = km.fit_predict(X)
ph_share = [float((pred[true==1]==c).mean()) if (pred==c).sum() else 0.0 for c in range(3)]
pc = int(np.argmax(ph_share))
yb = (pred == pc).astype(int)
prec, rec, f1, _ = precision_recall_fscore_support(true, yb, average='binary', pos_label=1)
out = dict(precision=round(prec,4), recall=round(rec,4), f1=round(f1,4),
           ari=round(adjusted_rand_score(true,pred),4), nmi=round(normalized_mutual_info_score(true,pred),4),
           ph_cluster=pc, ph_share=[round(x,4) for x in ph_share])
comp = []
for c in range(3):
    idx = np.where(pred==c)[0]
    comp.append(dict(cluster=int(c), size=int(len(idx)),
                     n_phishing=int((true[idx]==1).sum()),
                     n_spam=int(((kind=='垃圾') & (pred==c)).sum()),
                     n_normal=int(((kind=='正常') & (pred==c)).sum())))
out['comp'] = comp
out['ph_dist'] = [c['n_phishing'] for c in comp]
json.dump(out, open('hybrid_result_v3.json','w',encoding='utf-8'), ensure_ascii=False, indent=1)
print('seeded K=3:', json.dumps({k:v for k,v in out.items() if k!='comp'}, ensure_ascii=False))
for c in comp: print(f"  簇{c['cluster']}: 规模{c['size']} 钓鱼{c['n_phishing']} 垃圾{c['n_spam']} 正常{c['n_normal']}")
