# -*- coding: utf-8 -*-
"""钓鱼分数排序实验：按钓鱼信号特征加权打分，Top-N 命中率"""
import json, re
import numpy as np
import openpyxl

XLSX = 'data/法语语料库_最终版_钓鱼邮件319条+垃圾短信1272条.xlsx'
URL_RE = re.compile(r'https?://|www\.|\b[a-z0-9-]+\.(com|net|org|fr|eu|io|ly|xyz|top|club|site|info)\b', re.I)
SHORT_RE = re.compile(r'\b(bit\.ly|t\.co|goo\.gl|tinyurl|shorturl|ow\.ly|is\.gd|buff\.ly|cutt\.ly)\b|https?://[^\s]{1,12}\b', re.I)
MONEY_RE = re.compile(r'[€$£]|\b\d[\d\s.,]{2,}\s?(euros?|€|dollars?)\b', re.I)
BANK = re.compile(r'\b(banque|compte|relevé|identifiant|identif|code|mot de passe|carte|crédit|virement|transaction|sécurit|paypal|amazon|dhl|la poste|imp[ôô]ts|police|gendarmerie|arcep|free|orange|sosh|bouygues)\b', re.I)
THREAT = re.compile(r'\b(p[ée]nalit[ée]|amende|poursuites|pirat[ée]|ransomware|webcam|chantage|dossier|proc[èe]s|prison|virus|infection)\b', re.I)

def feats(t):
    tl = t.lower()
    return {'bank': 1 if BANK.search(tl) else 0,
            'url': 1 if URL_RE.search(tl) else 0,
            'threat': 1 if THREAT.search(tl) else 0,
            'money': 1 if MONEY_RE.search(tl) else 0,
            'short': 1 if SHORT_RE.search(tl) else 0,
            'len': len(t)/400.0}

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
rows = []
ws = wb['钓鱼邮件语料']
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not row[0] or not row[1]: continue
    rows.append({'id': str(row[0]), 'text': str(row[1]), 'true': 1, 'kind': '钓鱼'})
ws2 = wb['垃圾短信数据']
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0 or not row[0] or not row[1]: continue
    rows.append({'id': str(row[0]), 'text': str(row[1]), 'true': 0,
                 'kind': '垃圾' if int(row[2] if row[2] is not None else 1)==1 else '正常'})
wb.close()

# 打分：权重来自特征区分度实验（正信号加权）
W = {'bank': 2.0, 'url': 1.5, 'threat': 2.0, 'money': 1.0, 'short': 1.0, 'len': 0.8}
for r in rows:
    f = feats(r['text'])
    r['score'] = round(sum(W[k]*f[k] for k in W), 3)

rows.sort(key=lambda r: -r['score'])
total_ph = sum(1 for r in rows if r['true']==1)

# Top-N 命中率
N_list = [10, 20, 30, 50, 81, 100, 150, 200]
hits = []
for n in N_list:
    top = rows[:n]
    ph = sum(1 for r in top if r['true']==1)
    hits.append({'n': n, 'hit': ph, 'prec': round(ph/n, 3), 'recall': round(ph/total_ph, 3)})

# 分数分布对比（钓鱼 vs 非钓鱼的均分）
ph_scores = [r['score'] for r in rows if r['true']==1]
np_scores = [r['score'] for r in rows if r['true']==0]
out = {'total': len(rows), 'phishing': total_ph,
       'hits': hits,
       'mean_score': {'phishing': round(float(np.mean(ph_scores)),3), 'non_phishing': round(float(np.mean(np_scores)),3)},
       'top10_ids': [r['id'] for r in rows[:10]]}
json.dump(out, open('score_ranking.json','w',encoding='utf-8'), ensure_ascii=False, indent=1)
print('样本', out['total'], '钓鱼', out['phishing'])
print('钓鱼均分', out['mean_score']['phishing'], '/ 非钓鱼均分', out['mean_score']['non_phishing'])
for h in hits:
    print(f"  Top-{h['n']:>3}: 命中{h['hit']:>2} 精确率{h['prec']:.1%} 覆盖钓鱼{h['recall']:.1%}")
print('Top10 样本id:', out['top10_ids'])
